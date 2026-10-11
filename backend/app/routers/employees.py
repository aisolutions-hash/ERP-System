"""Employee Management (HR & Payroll Department) — CRUD, search & filters.

Sensitive personal identifiers (Aadhaar, PAN, bank account, IFSC, bank, branch)
are admin-only: full values are returned only to an admin, never placed in URLs,
audit details, or error messages. List responses always mask the identity and
bank-account numbers regardless of role.
"""
import csv
import io
import re
from datetime import date
from typing import Annotated

from fastapi import (
    APIRouter, Depends, File, HTTPException, Query, Response, UploadFile, status,
)
from fastapi.responses import StreamingResponse
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

from ..auth import CurrentUser, ManagerOrAdmin
from ..crud import get_or_404, write_audit
from ..database import get_db
from ..models import Employee, EmployeeStatus, UserRole
from ..schemas import EmployeeCreate, EmployeeOut, EmployeeUpdate
from ..services.employee_pdf import build_employees_pdf_bytes
from ..services.import_common import (
    build_column_map, cell_text, parse_date_value, read_table, row_to_dict,
)

router = APIRouter(prefix="/employees", tags=["employees"])

SENSITIVE_FIELDS = (
    "aadhaar_number", "pan_number", "bank_account_number",
    "ifsc_code", "bank_name", "branch_name",
)
MASKED_FIELDS = ("aadhaar_number", "pan_number", "bank_account_number")

# Indian format rules (applicable identifiers only; all fields optional).
AADHAAR_RE = re.compile(r"^[2-9]\d{11}$")
PAN_RE = re.compile(r"^[A-Z]{5}[0-9]{4}[A-Z]$")
IFSC_RE = re.compile(r"^[A-Z]{4}0[A-Z0-9]{6}$")


def _normalize_sensitive(data: dict) -> None:
    if data.get("aadhaar_number"):
        data["aadhaar_number"] = re.sub(r"[\s-]", "", data["aadhaar_number"])
    if data.get("bank_account_number"):
        data["bank_account_number"] = re.sub(r"\s", "", data["bank_account_number"])
    for field in ("pan_number", "ifsc_code"):
        if data.get(field):
            data[field] = data[field].strip().upper()


def _sensitive_errors(data: dict) -> list[str]:
    """Generic format-error messages (never echo the submitted value)."""
    errs = []
    aadhaar = data.get("aadhaar_number")
    if aadhaar and not AADHAAR_RE.match(aadhaar):
        errs.append("Invalid Aadhaar Number format")
    pan = data.get("pan_number")
    if pan and not PAN_RE.match(pan):
        errs.append("Invalid PAN Number format")
    ifsc = data.get("ifsc_code")
    if ifsc and not IFSC_RE.match(ifsc):
        errs.append("Invalid IFSC Code format")
    return errs


def _validate_sensitive(data: dict) -> None:
    errs = _sensitive_errors(data)
    if errs:
        raise HTTPException(status.HTTP_422_UNPROCESSABLE_ENTITY, errs[0])


def _mask(value: str | None) -> str:
    value = value or ""
    if len(value) <= 4:
        return "X" * len(value)
    return "X" * (len(value) - 4) + value[-4:]


def _serialize(emp: Employee, viewer, *, full: bool) -> dict:
    """Serialize an employee. `full=False` (list view) always masks identity and
    bank-account numbers and blanks the remaining sensitive fields; `full=True`
    returns full values only to an admin."""
    data = EmployeeOut.model_validate(emp).model_dump()
    if not (full and viewer.role == UserRole.admin):
        for field in MASKED_FIELDS:
            data[field] = _mask(data.get(field))
        for field in SENSITIVE_FIELDS:
            if field not in MASKED_FIELDS:
                data[field] = ""
    return data


# --- Export / import column metadata -------------------------------------
EXPORT_COLUMNS = [
    ("employee_code", "Employee ID"),
    ("name", "Name"),
    ("mobile", "Mobile"),
    ("email", "Email"),
    ("department", "Department"),
    ("designation", "Designation"),
    ("joining_date", "Joining Date"),
    ("status", "Status"),
    ("aadhaar_number", "Aadhaar Number"),
    ("pan_number", "PAN Number"),
    ("bank_account_number", "Bank Account Number"),
    ("ifsc_code", "IFSC Code"),
    ("bank_name", "Bank Name"),
    ("branch_name", "Branch Name"),
]
EXPORT_LABELS = dict(EXPORT_COLUMNS)
DEFAULT_EXPORT_FIELDS = [f for f, _ in EXPORT_COLUMNS if f not in SENSITIVE_FIELDS]

EMPLOYEE_HEADER_ALIASES = {
    "employee_code": ["Employee ID", "Employee Code", "Emp ID", "Code", "Employee Id"],
    "name": ["Name", "Employee Name", "Full Name"],
    "mobile": ["Mobile", "Mobile Number", "Phone", "Phone Number", "Contact"],
    "email": ["Email", "Email Address", "Mail"],
    "department": ["Department", "Dept"],
    "designation": ["Designation", "Role", "Position"],
    "joining_date": ["Joining Date", "Date of Joining", "DOJ", "Join Date"],
    "status": ["Status"],
    "aadhaar_number": ["Aadhaar Number", "Aadhaar", "Aadhar", "Aadhar Number"],
    "pan_number": ["PAN Number", "PAN", "Pan No"],
    "bank_account_number": ["Bank Account Number", "Bank Account", "Account Number",
                            "Account No", "Bank A/c", "Bank AC"],
    "ifsc_code": ["IFSC Code", "IFSC", "IFSC Code."],
    "bank_name": ["Bank Name", "Bank"],
    "branch_name": ["Branch Name", "Branch"],
}

MAX_IMPORT_BYTES = 5 * 1024 * 1024   # 5 MB
MAX_IMPORT_ROWS = 1000


def _mask_sensitive_copy(data: dict) -> dict:
    """Masked copy for previews: identity/bank numbers keep the last 4, the
    remaining sensitive fields are blanked."""
    out = dict(data)
    for field in MASKED_FIELDS:
        out[field] = _mask(out.get(field))
    for field in SENSITIVE_FIELDS:
        if field not in MASKED_FIELDS:
            out[field] = ""
    return out


def _parse_status(value: str) -> str | None:
    for st in EmployeeStatus:
        if st.value.lower() == value.strip().lower():
            return st.value
    return None


def _parse_import_rows(headers, rows, existing_codes: set[str]):
    """Validate and normalize uploaded rows. Returns (colmap, parsed_rows).

    Each parsed row: {row, data, errors, duplicate} where duplicate is
    'file' | 'database' | None. `data` holds normalized values ready to persist.
    """
    colmap = build_column_map(headers, EMPLOYEE_HEADER_ALIASES)
    parsed = []
    seen_codes: dict[str, int] = {}
    for i, raw in enumerate(rows):
        mapped = row_to_dict(colmap, raw)
        if all(cell_text(v) == "" for v in mapped.values()):
            continue
        row_no = i + 2  # +1 for header, +1 for 1-based
        data = {
            "employee_code": cell_text(mapped.get("employee_code")),
            "name": cell_text(mapped.get("name")),
            "mobile": cell_text(mapped.get("mobile")),
            "email": cell_text(mapped.get("email")),
            "department": cell_text(mapped.get("department")),
            "designation": cell_text(mapped.get("designation")),
            "status": EmployeeStatus.active.value,
            "status_provided": False,
            "joining_date": None,
        }
        for field in SENSITIVE_FIELDS:
            data[field] = cell_text(mapped.get(field))
        errors = []
        if not data["employee_code"]:
            errors.append("Employee ID is required")
        if not data["name"]:
            errors.append("Employee Name is required")

        raw_status = cell_text(mapped.get("status"))
        if raw_status:
            parsed_status = _parse_status(raw_status)
            if parsed_status is None:
                errors.append("Status must be 'Active' or 'Inactive'")
            else:
                data["status"] = parsed_status
                data["status_provided"] = True

        raw_date = cell_text(mapped.get("joining_date"))
        if raw_date:
            jd = parse_date_value(mapped.get("joining_date"))
            if jd is None:
                errors.append("Joining Date is not a valid date")
            data["joining_date"] = jd

        _normalize_sensitive(data)
        errors.extend(_sensitive_errors(data))

        duplicate = None
        code = data["employee_code"]
        if code:
            if code in seen_codes:
                duplicate = "file"
            elif code in existing_codes:
                duplicate = "database"
            seen_codes[code] = row_no

        parsed.append({"row": row_no, "data": data, "errors": errors,
                       "duplicate": duplicate})
    return colmap, parsed


def _ensure_upload_ok(content: bytes) -> None:
    if len(content) > MAX_IMPORT_BYTES:
        raise HTTPException(status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                            f"File too large (limit {MAX_IMPORT_BYTES // (1024 * 1024)} MB)")


def _read_upload(filename: str, content: bytes):
    _ensure_upload_ok(content)
    try:
        headers, rows = read_table(filename or "upload.csv", content)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if not rows:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "No data rows found in file")
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            f"Too many rows (limit {MAX_IMPORT_ROWS})")
    return headers, rows


def _existing_code_map(db: Session) -> dict[str, Employee]:
    return {e.employee_code: e for e in db.scalars(select(Employee)).all()}


def _import_fields(data: dict, user) -> dict:
    """Employee column values from an imported row, honoring the admin-only
    rule for sensitive fields (non-admins never set them)."""
    fields = {
        "employee_code": data["employee_code"],
        "name": data["name"],
        "mobile": data["mobile"],
        "email": data["email"],
        "department": data["department"],
        "designation": data["designation"],
        "joining_date": data["joining_date"],
        "status": EmployeeStatus(data["status"]),
    }
    for field in SENSITIVE_FIELDS:
        fields[field] = data[field] if user.role == UserRole.admin else ""
    return fields


def _csv_response(headers, rows, filename):
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(headers)
    w.writerows(rows)
    return StreamingResponse(
        io.BytesIO(buf.getvalue().encode("utf-8-sig")),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _excel_response(headers, rows, filename):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(headers)
    fill = PatternFill("solid", fgColor="1E3A8A")
    font = Font(color="FFFFFF", bold=True)
    for c in ws[1]:
        c.fill = fill
        c.font = font
    for r in rows:
        ws.append(r)
    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A"].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


def _resolve_export_fields(columns: str, include_sensitive: bool, viewer) -> list[str]:
    requested = [c.strip() for c in (columns or "").split(",") if c.strip()]
    if requested:
        fields = [c for c in requested if c in EXPORT_LABELS]
    else:
        fields = list(DEFAULT_EXPORT_FIELDS)
        if include_sensitive:
            fields += [f for f, _ in EXPORT_COLUMNS if f in SENSITIVE_FIELDS]
    allowed_sensitive = include_sensitive and viewer.role == UserRole.admin
    if not allowed_sensitive:
        fields = [f for f in fields if f not in SENSITIVE_FIELDS]
    return fields or list(DEFAULT_EXPORT_FIELDS)


def _export_value(field: str, value, *, unmask: bool):
    if field not in SENSITIVE_FIELDS:
        return value if value is not None else ""
    value = value or ""
    if unmask:
        return value
    if not value:
        return ""
    if field in MASKED_FIELDS:
        return _mask(value)
    return "\u2022" * min(len(str(value)), 8)


@router.get("", response_model=dict)
def list_employees(
    db: Annotated[Session, Depends(get_db)],
    viewer: CurrentUser,
    search: str = "",
    department: str = "",
    designation: str = "",
    status_filter: str = Query("", alias="status"),
    page: int = Query(1, ge=1),
    page_size: int = Query(25, ge=1, le=500),
):
    stmt = select(Employee)
    if search:
        like = f"%{search}%"
        stmt = stmt.where(or_(
            Employee.employee_code.ilike(like), Employee.name.ilike(like),
            Employee.mobile.ilike(like), Employee.email.ilike(like),
        ))
    if department:
        stmt = stmt.where(Employee.department == department)
    if designation:
        stmt = stmt.where(Employee.designation == designation)
    if status_filter:
        stmt = stmt.where(Employee.status == status_filter)
    total = db.scalar(select(func.count()).select_from(stmt.subquery()))
    rows = db.scalars(
        stmt.order_by(Employee.employee_code).offset((page - 1) * page_size).limit(page_size)
    ).all()
    return {"items": [_serialize(e, viewer, full=False) for e in rows],
            "total": total, "page": page, "page_size": page_size}


@router.get("/filters", response_model=dict)
def employee_filters(db: Annotated[Session, Depends(get_db)], _: CurrentUser):
    depts = db.scalars(select(Employee.department).where(Employee.department != "")
                       .distinct().order_by(Employee.department)).all()
    desigs = db.scalars(select(Employee.designation).where(Employee.designation != "")
                        .distinct().order_by(Employee.designation)).all()
    return {"departments": list(depts), "designations": list(desigs),
            "statuses": [s.value for s in EmployeeStatus]}


@router.get("/import/template")
def import_template():
    """Downloadable CSV template with the correct column headers."""
    headers = [label for _, label in EXPORT_COLUMNS]
    sample = ["EMP001", "Jane Doe", "9876543210", "jane@example.com",
              "Finance", "Accountant", "2025-04-01", "Active",
              "234567890123", "ABCDE1234F", "123456789012",
              "HDFC0001234", "HDFC Bank", "Main Branch"]
    return _csv_response(headers, [sample], "employee_import_template.csv")


@router.post("/import/preview", response_model=dict)
async def preview_employees_import(
    db: Annotated[Session, Depends(get_db)],
    user: ManagerOrAdmin,
    file: UploadFile = File(...),
):
    """Parse + validate an upload WITHOUT saving. Sensitive values are masked."""
    content = await file.read()
    headers, rows = _read_upload(file.filename or "", content)
    colmap, parsed = _parse_import_rows(headers, rows, set(_existing_code_map(db)))
    if "employee_code" not in colmap.values() and "name" not in colmap.values():
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "No recognized columns found — the file needs at least "
                            "an Employee ID or Name column")
    preview_rows = [{
        "row": p["row"],
        "status": "error" if p["errors"] else ("duplicate" if p["duplicate"] else "valid"),
        "errors": p["errors"],
        "data": _mask_sensitive_copy(p["data"]),
    } for p in parsed]
    summary = {
        "total": len(parsed),
        "valid": sum(1 for r in preview_rows if r["status"] == "valid"),
        "errors": sum(1 for r in preview_rows if r["status"] == "error"),
        "duplicates": sum(1 for r in preview_rows if r["status"] == "duplicate"),
    }
    return {
        "filename": file.filename or "",
        "headers": headers,
        "mapped_columns": {headers[idx]: canon for idx, canon in colmap.items()},
        "detected_fields": sorted(set(colmap.values())),
        "columns": [{"field": f, "label": label, "sensitive": f in SENSITIVE_FIELDS}
                    for f, label in EXPORT_COLUMNS],
        "rows": preview_rows,
        "summary": summary,
    }


@router.post("/import", response_model=dict)
async def import_employees(
    db: Annotated[Session, Depends(get_db)],
    user: ManagerOrAdmin,
    file: UploadFile = File(...),
    duplicate_action: str = Query("skip", pattern="^(skip|update)$",
                                  description="Action for existing Employee IDs"),
):
    """Bulk-import employees from CSV/Excel. Valid rows are saved; invalid rows
    are reported. Duplicate Employee IDs are skipped (or updated) — existing
    records are never silently overwritten."""
    content = await file.read()
    headers, rows = _read_upload(file.filename or "", content)
    existing = _existing_code_map(db)
    colmap, parsed = _parse_import_rows(headers, rows, set(existing))
    if "employee_code" not in colmap.values() and "name" not in colmap.values():
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "No recognized columns found in file")

    imported = updated = skipped = 0
    errors: list[dict] = []
    for p in parsed:
        data = p["data"]
        if p["errors"]:
            errors.append({"row": p["row"], "employee_code": data["employee_code"],
                           "message": "; ".join(p["errors"])})
            continue
        if p["duplicate"] == "file":
            skipped += 1
            continue
        current = existing.get(data["employee_code"])
        if current is not None:
            if duplicate_action == "update":
                # Only apply columns that were actually provided with a value —
                # blank/absent cells never clear existing data.
                fields = _import_fields(data, user)
                if not data.get("status_provided"):
                    fields.pop("status", None)
                for key, value in fields.items():
                    if value in (None, ""):
                        continue
                    setattr(current, key, value)
                updated += 1
            else:
                skipped += 1
            continue
        emp = Employee(**_import_fields(data, user))
        db.add(emp)
        existing[data["employee_code"]] = emp
        imported += 1

    db.commit()
    write_audit(db, user, "IMPORT", "employees", None,
                f"Employee import: {imported} created, {updated} updated, "
                f"{skipped} skipped, {len(errors)} row errors")
    return {
        "filename": file.filename or "",
        "summary": {"total": len(parsed), "imported": imported, "updated": updated,
                    "skipped": skipped, "errors": len(errors)},
        "errors": errors,
    }


@router.get("/export")
def export_employees(
    db: Annotated[Session, Depends(get_db)],
    viewer: CurrentUser,
    fmt: str = Query("csv", alias="format"),
    search: str = "",
    department: str = "",
    designation: str = "",
    status_filter: str = Query("", alias="status"),
    ids: str = "",
    columns: str = "",
    include_sensitive: bool = False,
    unmask_sensitive: bool = False,
):
    """Export employees to CSV / Excel / PDF.

    Sensitive fields are excluded unless `include_sensitive=true` AND the caller
    is an admin; when included they are masked unless `unmask_sensitive=true`
    (admin only). Selecting specific records via `ids` overrides the filters.
    """
    fmt = (fmt or "csv").lower()
    if fmt == "excel":
        fmt = "xlsx"
    if fmt not in ("csv", "xlsx", "pdf"):
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "format must be one of csv, xlsx, pdf")
    if include_sensitive and viewer.role != UserRole.admin:
        raise HTTPException(status.HTTP_403_FORBIDDEN,
                            "Not permitted to export sensitive fields")

    fields = _resolve_export_fields(columns, include_sensitive, viewer)
    labels = [EXPORT_LABELS[f] for f in fields]
    unmask = bool(unmask_sensitive and include_sensitive and viewer.role == UserRole.admin)

    id_list = [int(x) for x in ids.split(",") if x.strip().isdigit()]
    stmt = select(Employee)
    if id_list:
        stmt = stmt.where(Employee.id.in_(id_list))
    else:
        if search:
            like = f"%{search}%"
            stmt = stmt.where(or_(
                Employee.employee_code.ilike(like), Employee.name.ilike(like),
                Employee.mobile.ilike(like), Employee.email.ilike(like),
            ))
        if department:
            stmt = stmt.where(Employee.department == department)
        if designation:
            stmt = stmt.where(Employee.designation == designation)
        if status_filter:
            stmt = stmt.where(Employee.status == status_filter)
    emps = db.scalars(stmt.order_by(Employee.employee_code)).all()

    rows = []
    for e in emps:
        d = EmployeeOut.model_validate(e).model_dump()
        rows.append([_export_value(f, d.get(f), unmask=unmask) for f in fields])

    stamp = date.today().isoformat()
    if id_list:
        subtitle = f"Selected employees: {len(emps)}"
    else:
        subtitle = f"Filtered list — {len(emps)} record(s)"
    if include_sensitive and viewer.role == UserRole.admin:
        subtitle += "  |  sensitive fields masked" if not unmask else "  |  sensitive fields included"

    if fmt == "csv":
        return _csv_response(labels, rows, f"employees_{stamp}.csv")
    if fmt == "xlsx":
        return _excel_response(labels, rows, f"employees_{stamp}.xlsx")
    pdf = build_employees_pdf_bytes(labels, rows, title="Employee List",
                                    subtitle=subtitle)
    return Response(
        content=pdf,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename=employees_{stamp}.pdf"},
    )


@router.post("", response_model=EmployeeOut, status_code=status.HTTP_201_CREATED)
def create_employee(body: EmployeeCreate, db: Annotated[Session, Depends(get_db)],
                    user: ManagerOrAdmin):
    if db.query(Employee).filter(Employee.employee_code == body.employee_code).first():
        raise HTTPException(status.HTTP_409_CONFLICT, "Employee ID already exists")
    data = body.model_dump()
    if user.role == UserRole.admin:
        _normalize_sensitive(data)
        _validate_sensitive(data)
    else:
        for field in SENSITIVE_FIELDS:
            data[field] = ""
    emp = Employee(**data)
    db.add(emp)
    db.commit()
    db.refresh(emp)
    write_audit(db, user, "CREATE", "employees", emp.id, f"Created employee {emp.name}")
    return _serialize(emp, user, full=True)


@router.get("/{employee_id}", response_model=EmployeeOut)
def get_employee(employee_id: int, db: Annotated[Session, Depends(get_db)],
                 viewer: CurrentUser):
    return _serialize(get_or_404(db, Employee, employee_id), viewer, full=True)


@router.patch("/{employee_id}", response_model=EmployeeOut)
def update_employee(employee_id: int, body: EmployeeUpdate,
                    db: Annotated[Session, Depends(get_db)], user: ManagerOrAdmin):
    emp = get_or_404(db, Employee, employee_id)
    data = body.model_dump(exclude_unset=True)
    code = data.get("employee_code")
    if code and code != emp.employee_code and \
            db.query(Employee).filter(Employee.employee_code == code).first():
        raise HTTPException(status.HTTP_409_CONFLICT, "Employee ID already exists")
    if user.role == UserRole.admin:
        _normalize_sensitive(data)
        _validate_sensitive(data)
    else:
        # Non-admins cannot read or modify sensitive fields.
        for field in SENSITIVE_FIELDS:
            data.pop(field, None)
    for key, value in data.items():
        if value is not None and hasattr(emp, key):
            setattr(emp, key, value)
    db.commit()
    db.refresh(emp)
    write_audit(db, user, "UPDATE", "employees", emp.id, f"Updated employee {emp.name}")
    return _serialize(emp, user, full=True)


@router.delete("/{employee_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_employee(employee_id: int, db: Annotated[Session, Depends(get_db)],
                    user: ManagerOrAdmin):
    emp = get_or_404(db, Employee, employee_id)
    db.delete(emp)
    db.commit()
    write_audit(db, user, "DELETE", "employees", employee_id, f"Deleted employee {emp.name}")
    return Response(status_code=status.HTTP_204_NO_CONTENT)
