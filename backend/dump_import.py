from openpyxl import load_workbook
wb = load_workbook(r'C:\Users\user\Downloads\IMPORT PO FORMAT Dung guan Hobiba technology pvt ltd. (5).xlsx')
ws = wb.active
for r in range(1, 45):
    row = []
    for c in range(1, 12):
        v = ws.cell(r,c).value
        if v is None:
            v=''
        row.append(str(v).replace('\n',' '))
    print(f'{r}: ' + ' | '.join(row))
